"""O'quv dasturi: taqvim-mavzu reja va ko'nikmalar xaritasi (FR-04, FR-05).

Standart reja — "AHA! Matematika 5-sinf" darsligi bo'yicha 2026–2027 o'quv yili taqvim-mavzu rejasi (aha_plan.py).
O'qituvchi o'z rejasini (PDF/DOCX/XLSX) yuklasa, AI mavzularga ajratadi va shu reja ishlatiladi.
"""
import io
import re
import zipfile

from sqlalchemy import delete, select

from . import aha_plan, clock, db, llm, service
from .curriculum import SKILL_NAMES, SKILLS
from .models import CurriculumTopic, Lesson
from .problems import TEMPLATES


def default_rows():
    return aha_plan.grouped()


def assign_topics(s, class_id):
    """Darslarni tartib raqami (seq) bo'yicha reja mavzulariga bog'laydi; o'tkazilgan darslarning mavzusi o'zgarmaydi."""
    topics = s.scalars(select(CurriculumTopic).where(CurriculumTopic.class_id == class_id).order_by(CurriculumTopic.order)).all()
    slots = [t for t in topics for _ in range(max(1, t.hours))]
    for lesson in s.scalars(select(Lesson).where(Lesson.class_id == class_id).order_by(Lesson.date)):
        if lesson.conducted_at is not None or lesson.seq is None:
            continue
        t = slots[lesson.seq] if lesson.seq < len(slots) else None
        lesson.curriculum_topic_id = t.id if t else None
        lesson.topic = t.title if t else "Takrorlash va mustahkamlash"


def ensure_default(class_id=None):
    with db.session() as s:
        cls = service.get_class(s, class_id)
        if s.scalars(select(CurriculumTopic).where(CurriculumTopic.class_id == cls.id)).first():
            return
        for i, row in enumerate(default_rows(), start=1):
            s.add(CurriculumTopic(class_id=cls.id, order=i, week=row["week"], title=row["title"], hours=row["hours"],
                                  skills=row["skills"], prerequisites=row["prerequisites"], template=row["template"],
                                  source="standart", quarter=row["quarter"], chapter=row["chapter"],
                                  chapter_title=row["chapter_title"], kind=row["kind"], points=row["points"],
                                  textbook=row["textbook"], workbook=row["workbook"], lesson_no=row["lesson_no"]))


def topic_refs(topic: CurriculumTopic | None) -> dict:
    if not topic:
        return {}
    return {k: v for k, v in {"textbook": topic.textbook, "workbook": topic.workbook, "chapter": topic.chapter_title,
                              "lesson_no": topic.lesson_no, "kind": topic.kind, "points": topic.points}.items() if v}


def view(class_id=None):
    ensure_default(class_id)
    with db.session() as s:
        cls = service.get_class(s, class_id)
        rows = s.scalars(select(CurriculumTopic).where(CurriculumTopic.class_id == cls.id).order_by(CurriculumTopic.order)).all()
        done = {t for (t,) in s.execute(select(Lesson.curriculum_topic_id).where(Lesson.class_id == cls.id, Lesson.conducted_at.is_not(None)))}
        current = s.scalars(select(Lesson).where(Lesson.class_id == cls.id, Lesson.date >= clock.today()).order_by(Lesson.date)).first()
        return {
            "class": cls.name,
            "class_id": cls.id,
            "skills": [{"key": k, "name": n} for k, n in SKILLS],
            "templates": [{"key": k, "title": v} for k, v in TEMPLATES.items()],
            "calendar": aha_plan.calendar_view(clock.today()),
            "current_topic_id": current.curriculum_topic_id if current else None,
            "topics": [{"id": r.id, "order": r.order, "week": r.week, "title": r.title, "hours": r.hours, "passed": r.id in done,
                        "skills": r.skills, "prerequisites": r.prerequisites, "template": r.template, "source": r.source,
                        "quarter": r.quarter, "chapter": r.chapter, "chapter_title": r.chapter_title, "kind": r.kind,
                        "points": r.points, "textbook": r.textbook, "workbook": r.workbook, "lesson_no": r.lesson_no}
                       for r in rows],
        }


def extract_text(data: bytes, filename: str) -> str:
    name = filename.lower()
    if name.endswith(".docx"):
        xml = zipfile.ZipFile(io.BytesIO(data)).read("word/document.xml").decode("utf8", errors="ignore")
        paras = re.findall(r"<w:p[ >].*?</w:p>", xml, flags=re.S)
        return "\n".join("".join(re.findall(r"<w:t[^>]*>(.*?)</w:t>", p)) for p in paras)
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True)
        lines = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
                if cells:
                    lines.append(" | ".join(cells))
        return "\n".join(lines)
    if name.endswith(".pdf"):
        import pymupdf
        doc = pymupdf.open(stream=data, filetype="pdf")
        return "\n".join(page.get_text() for page in doc)
    if name.endswith((".txt", ".csv")):
        return data.decode("utf-8", errors="ignore")
    raise ValueError("Hozircha PDF, DOCX, XLSX yoki TXT fayl qabul qilinadi (rasm uchun OCR keyingi versiyada)")


def _fallback_split(text: str) -> list:
    topics = []
    for line in text.splitlines():
        clean = re.sub(r"^\s*\d+[.)]\s*", "", line).strip(" |\t")
        if len(clean) < 6 or not any(ch.isalpha() for ch in clean):
            continue
        hours = re.search(r"(\d+)\s*soat", clean)
        title = re.sub(r"\|?\s*\d+\s*soat.*$", "", clean).strip(" |")
        topics.append({"title": title[:200], "week": None, "hours": int(hours.group(1)) if hours else 1, "skills": [], "prerequisites": []})
    return topics[:40]


def import_preview(data: bytes, filename: str) -> dict:
    text = extract_text(data, filename)
    if len(text.strip()) < 10:
        raise ValueError("Fayldan matn topilmadi")
    items = llm.split_curriculum(text, SKILLS)
    source = "gpt"
    if not items:
        items, source = _fallback_split(text), "shablon"
    topics = []
    for it in items:
        if not isinstance(it, dict) or not str(it.get("title", "")).strip():
            continue
        topics.append({
            "title": str(it["title"]).strip()[:200],
            "week": it.get("week") if isinstance(it.get("week"), int) else None,
            "hours": max(1, min(20, int(it.get("hours") or 1))) if str(it.get("hours") or "1").isdigit() else 1,
            "skills": [k for k in it.get("skills") or [] if k in SKILL_NAMES],
            "prerequisites": [k for k in it.get("prerequisites") or [] if k in SKILL_NAMES],
        })
    if not topics:
        raise ValueError("Mavzular ajratilmadi")
    return {"topics": topics, "source": source, "chars": len(text)}


def _clear(s, class_id):
    for lesson in s.scalars(select(Lesson).where(Lesson.class_id == class_id)):
        lesson.curriculum_topic_id = None
    s.flush()
    s.execute(delete(CurriculumTopic).where(CurriculumTopic.class_id == class_id))


def confirm(topics: list, class_id=None):
    if not topics:
        raise ValueError("Mavzular ro'yxati bo'sh")
    with db.session() as s:
        cls = service.get_class(s, class_id)
        _clear(s, cls.id)
        lesson_no = 1
        for i, t in enumerate(topics, start=1):
            title = str(t.get("title", "")).strip()[:200]
            if not title:
                continue
            hours = int(t.get("hours") or 1)
            template = next((k for k, v in TEMPLATES.items() if k.replace("_", " ") in title.lower() or v.lower()[:18] in title.lower()), None)
            s.add(CurriculumTopic(class_id=cls.id, order=i, week=t.get("week"), title=title, hours=hours,
                                  skills=[k for k in t.get("skills") or [] if k in SKILL_NAMES],
                                  prerequisites=[k for k in t.get("prerequisites") or [] if k in SKILL_NAMES],
                                  template=template, source="yuklangan", kind="dars", lesson_no=lesson_no))
            lesson_no += hours
        s.flush()
        assign_topics(s, cls.id)
    return view(cls.id)


def reset_default(class_id=None):
    with db.session() as s:
        cls = service.get_class(s, class_id)
        _clear(s, cls.id)
        cid = cls.id
    ensure_default(cid)
    with db.session() as s:
        assign_topics(s, cid)
    return view(cid)
