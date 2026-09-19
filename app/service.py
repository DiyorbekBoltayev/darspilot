"""Biznes-mantiq: diagnostika, skanerlash, baholash, dars ssenariysi, panel. API uchun oddiy dict qaytaradi."""
import io
import random
import secrets
import time
from concurrent.futures import ThreadPoolExecutor
from datetime import date, datetime, timedelta

import cv2
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import delete, func, select

from . import clock, db, grading, llm, omr, pdfgen, problems, scenario, simulate
from . import methods as methods_mod
from .config import ATTENTION_GAP_ALERT, MINUTES_SAVED_PER_HOMEWORK, MINUTES_SAVED_PER_STUDENT
from .curriculum import ERROR_ACTIONS, ERRORS, LEVEL_NAMES, LEVELS, QUESTION_STEP, SKILLS, level_for_score
from .methods import METHODS, STAGES
from .models import (Attention, ClassGroups, CurriculumTopic, CustomMethod, DiagAssign, Diagnostic, DiagProblem, DiagSummary, Grade,
                     Homework, Lesson, LessonPlan, LlmCall, MethodPreference, Response, Result, Scan, SchoolClass, Setting,
                     SkillScore, Student)
from .storage import storage


class NotFound(Exception):
    pass


# Uy vazifasi suratlari fonda tahlil qilinadi: o'qituvchi sinfda yurib suratga olaveradi, kutmaydi.
_BG = ThreadPoolExecutor(max_workers=3, thread_name_prefix="darspilot-bg")


def avg(xs):
    xs = [x for x in xs if x is not None]
    return sum(xs) / len(xs) if xs else 0.0


def fmt_date(d: date) -> str:
    return d.strftime("%d.%m.%Y")


# ------------------------------------------------------------------ sozlamalar
def gap_alert() -> int:
    """"N dars e'tiborsiz" chegarasi (FR-19), standart 5."""
    with db.session() as s:
        row = s.get(Setting, "gap_alert")
        return int(row.value) if row else ATTENTION_GAP_ALERT


def settings_view():
    return {"gap_alert": gap_alert()}


def save_settings(gap: int):
    if not 2 <= gap <= 15:
        raise ValueError("Chegara 2 va 15 orasida bo'lishi kerak")
    with db.session() as s:
        row = s.get(Setting, "gap_alert")
        if row:
            row.value = str(gap)
        else:
            s.add(Setting(key="gap_alert", value=str(gap)))


# ------------------------------------------------------------------ umumiy so'rovlar
def get_class(s, class_id=None):
    if class_id:
        cls = s.get(SchoolClass, class_id)
        if not cls:
            raise NotFound
        return cls
    return s.scalars(select(SchoolClass).order_by(SchoolClass.name)).first()


def class_students(s, class_id):
    return s.scalars(select(Student).where(Student.class_id == class_id).order_by(Student.journal_no)).all()


def skill_map(s, class_id):
    out = {}
    rows = s.scalars(select(SkillScore).join(Student, Student.id == SkillScore.student_id).where(Student.class_id == class_id))
    for sc in rows:
        out.setdefault(sc.student_id, {})[sc.skill] = sc.score
    return out


def today_lesson(s, class_id, topic=None):
    """Bugungi dars; bugun dars bo'lmasa — oxirgi o'tgan dars (e'tibor jurnali unga yoziladi)."""
    lesson = s.scalars(select(Lesson).where(Lesson.class_id == class_id, Lesson.date == clock.today())).first()
    if not lesson:
        lesson = s.scalars(select(Lesson).where(Lesson.class_id == class_id, Lesson.date < clock.today())
                           .order_by(Lesson.date.desc())).first()
    if not lesson:
        lesson = Lesson(class_id=class_id, date=clock.today(), topic=topic, seq=0)
        s.add(lesson)
        s.flush()
    return lesson


def lesson_template(s, lesson):
    """Dars mavzusiga mos diagnostik shablon (bo'lmasa None)."""
    if lesson.curriculum_topic_id:
        topic = s.get(CurriculumTopic, lesson.curriculum_topic_id)
        if topic and topic.template:
            return topic.template
        # BSB/ChSB o'zi mavzu emas — u endigina o'tilgan bo'limni tekshiradi, shuning uchun shabloni ham o'shanikidir
        if topic and topic.kind in ("bsb", "chsb"):
            prev = s.scalars(select(CurriculumTopic)
                             .where(CurriculumTopic.class_id == topic.class_id, CurriculumTopic.order < topic.order,
                                    CurriculumTopic.template.is_not(None))
                             .order_by(CurriculumTopic.order.desc())).first()
            if prev:
                return prev.template
    text = (lesson.topic or "").lower()
    if "qarama" in text:
        return "qarama_qarshi"
    if "quvib" in text or "bir yo'nalish" in text:
        return "quvib_yetish"
    return None


def attention_gaps(s, class_id):
    """O'quvchi bilan oxirgi marta ishlangandan beri o'tgan darslar soni."""
    lesson_dates = s.scalars(select(Lesson.date).where(Lesson.class_id == class_id, Lesson.conducted_at.is_not(None))
                             .order_by(Lesson.date)).all()
    last = dict(s.execute(select(Attention.student_id, func.max(Lesson.date)).join(Lesson, Lesson.id == Attention.lesson_id)
                          .where(Lesson.class_id == class_id).group_by(Attention.student_id)).all())
    return {st.id: sum(1 for d in lesson_dates if last.get(st.id) is None or d > last[st.id])
            for st in class_students(s, class_id)}


def student_brief(st: Student):
    return {"id": st.id, "code": st.code, "name": st.full_name, "journal_no": st.journal_no}


def storage_url(key: str) -> str:
    return f"/api/files/{key}"


# ------------------------------------------------------------------ panel
def overview(class_id=None):
    N = gap_alert()
    with db.session() as s:
        cls = get_class(s, class_id)
        sts = class_students(s, cls.id)
        smap = skill_map(s, cls.id)
        gaps = attention_gaps(s, cls.id)
        diags = s.scalars(select(Diagnostic).where(Diagnostic.class_id == cls.id).order_by(Diagnostic.id.desc())).all()
        graded_count = s.scalar(select(func.count()).select_from(Result).join(Diagnostic, Diagnostic.id == Result.diagnostic_id)
                                .where(Diagnostic.class_id == cls.id)) or 0
        last_summary = s.execute(select(DiagSummary, Diagnostic).join(Diagnostic, Diagnostic.id == DiagSummary.diagnostic_id)
                                 .where(Diagnostic.class_id == cls.id).order_by(Diagnostic.id.desc())).first()
        last_plan = s.scalars(select(LessonPlan).where(LessonPlan.class_id == cls.id).order_by(LessonPlan.id.desc())).first()
        students = []
        for st in sts:
            sk = smap.get(st.id, {})
            a = avg(sk.values())
            students.append({**student_brief(st), "skills": {k: sk.get(k) for k, _ in SKILLS}, "avg": round(a, 3),
                             "gap": gaps[st.id], "alert": gaps[st.id] >= N,
                             "level": level_for_score(a) if sk else "B2"})
        return {
            "class": {"id": cls.id, "name": cls.name, "subject": cls.subject, "grade": cls.grade},
            "skills": [{"key": k, "name": n} for k, n in SKILLS],
            "students": students,
            "kpi": {"students": len(sts), "avg": round(100 * avg(x["avg"] for x in students)),
                    "neglected": sum(1 for x in students if x["alert"]),
                    "minutes_saved": round(graded_count * MINUTES_SAVED_PER_STUDENT)},
            "diagnostics": [{"id": d.id, "title": d.title, "date": fmt_date(d.date), "status": d.status} for d in diags],
            "summary": ({"diagnostic_id": last_summary[1].id, "title": last_summary[1].title, "source": last_summary[0].source,
                         **last_summary[0].summary["summary"]} if last_summary else None),
            "last_plan": ({"id": last_plan.id, "topic": last_plan.topic,
                           "created_at": last_plan.created_at.strftime("%d.%m.%Y %H:%M")} if last_plan else None),
            "gap_alert": N,
        }


def student_detail(sid):
    N = gap_alert()
    with db.session() as s:
        st = s.get(Student, sid)
        if not st:
            raise NotFound
        sk = {x.skill: x.score for x in s.scalars(select(SkillScore).where(SkillScore.student_id == sid))}
        gaps = attention_gaps(s, st.class_id)
        history = s.execute(select(Result, Diagnostic, DiagAssign.level).join(Diagnostic, Diagnostic.id == Result.diagnostic_id)
                            .join(DiagAssign, (DiagAssign.diagnostic_id == Result.diagnostic_id) & (DiagAssign.student_id == Result.student_id))
                            .where(Result.student_id == sid).order_by(Diagnostic.id.desc())).all()
        att = s.execute(select(Lesson.date, Lesson.topic).join(Attention, Attention.lesson_id == Lesson.id)
                        .where(Attention.student_id == sid).order_by(Lesson.date.desc())).all()
        return {
            **student_brief(st),
            "skills": [{"key": k, "name": n, "score": sk.get(k)} for k, n in SKILLS],
            "avg": round(avg(sk.values()), 3), "gap": gaps[sid], "gap_alert": N,
            "history": [{"diagnostic_id": d.id, "title": d.title, "date": fmt_date(d.date), "level": lvl,
                         "correct": r.correct, "total": r.total, "primary_text": r.details["primary_text"],
                         "steps": [{"step": x["step"], "ok": x["ok"]} for x in r.details["steps"]],
                         "feedback": {"student": r.feedback_student, "parent": r.feedback_parent, "teacher": r.feedback_teacher}}
                        for r, d, lvl in history],
            "attention": [{"date": fmt_date(d), "topic": t} for d, t in att],
        }


# ------------------------------------------------------------------ diagnostika
def list_diagnostics(class_id=None):
    with db.session() as s:
        q = select(Diagnostic, func.count(Response.student_id)).outerjoin(Response, Response.diagnostic_id == Diagnostic.id)
        if class_id:
            q = q.where(Diagnostic.class_id == class_id)
        rows = s.execute(q.group_by(Diagnostic.id).order_by(Diagnostic.id.desc())).all()
        return [{"id": d.id, "title": d.title, "template": d.template, "date": fmt_date(d.date), "status": d.status,
                 "lesson_id": d.lesson_id, "class_id": d.class_id, "responses": n} for d, n in rows]


def create_diagnostic(template: str | None = None, lesson_id: int | None = None, class_id: int | None = None,
                      use_llm: bool = True) -> int:
    with db.session() as s:
        if lesson_id:
            lesson = s.get(Lesson, lesson_id)
            if not lesson:
                raise NotFound
        else:
            lesson = today_lesson(s, get_class(s, class_id).id)
        template = template or lesson_template(s, lesson)
        lid, cid, ldate = lesson.id, lesson.class_id, lesson.date
        # Summativ kun (BSB/ChSB) bo'lsa — xuddi shu dvigatel bilan, lekin baho turi va maksimal bali boshqa
        topic_row = s.get(CurriculumTopic, lesson.curriculum_topic_id) if lesson.curriculum_topic_id else None
        kind = topic_row.kind if topic_row and topic_row.kind in ("bsb", "chsb") else "formativ"
        max_points = (topic_row.points if topic_row and topic_row.points else GRADE_KINDS[kind]) if kind != "formativ" else 10
        kind_title = topic_row.title if topic_row else None
    if template not in problems.TEMPLATES:
        raise ValueError("Bu mavzu uchun diagnostik shablon hali yo'q")
    seed = int(time.time() * 1000) % 10 ** 9
    specs = {lvl: problems.generate(template, lvl, seed) for lvl in LEVELS}
    if use_llm:
        llm.enrich_problems(specs)  # GPT chaqiruvlari tranzaksiyadan tashqarida
    with db.session() as s:
        title = problems.TEMPLATES[template]
        if kind != "formativ":
            title = f"{kind.upper()} · {kind_title or title}"[:200]
        d = Diagnostic(class_id=cid, lesson_id=lid, template=template, title=title, date=ldate,
                       kind=kind, max_points=max_points)
        s.add(d)
        s.flush()
        for lvl, spec in specs.items():
            s.add(DiagProblem(diagnostic_id=d.id, level=lvl, spec=spec))
        smap = skill_map(s, cid)
        for st in class_students(s, cid):
            sk = smap.get(st.id)
            s.add(DiagAssign(diagnostic_id=d.id, student_id=st.id, level=level_for_score(avg(sk.values())) if sk else "B2"))
        s.get(Lesson, lid).diagnostic_day = True
        did = d.id
    if use_llm:
        build_pdfs(did)
    return did


def _load(s, did):
    d = s.get(Diagnostic, did)
    if not d:
        raise NotFound
    specs = {p.level: p.spec for p in s.scalars(select(DiagProblem).where(DiagProblem.diagnostic_id == did))}
    assign = {a.student_id: a.level for a in s.scalars(select(DiagAssign).where(DiagAssign.diagnostic_id == did))}
    return d, specs, assign


def pdf_key(did, kind):
    return f"pdf/diag{did}_{kind}.pdf"


def last_feedback_map(s, class_id, before_id):
    """Har o'quvchining oxirgi diagnostikadagi shaxsiy feedbacki — keyingi kartochkasida qaytarish uchun."""
    rows = s.execute(select(Result.student_id, Result.feedback_student, Diagnostic.id)
                     .join(Diagnostic, Diagnostic.id == Result.diagnostic_id)
                     .where(Diagnostic.class_id == class_id, Diagnostic.id < before_id,
                            Result.feedback_student.is_not(None))
                     .order_by(Diagnostic.id)).all()
    return {sid: text for sid, text, _ in rows}


def diagnostic_cards(s, did):
    """Har bir o'quvchi uchun kartochka ma'lumoti: ism, kod, variant, masala va o'tgan ishdan feedback."""
    d, specs, assign = _load(s, did)
    fb = last_feedback_map(s, d.class_id, did)
    return [{"journal_no": st.journal_no, "code": st.code, "name": st.full_name,
             "level": assign[st.id], "spec": specs[assign[st.id]], "feedback": fb.get(st.id),
             "kind": d.kind, "max_points": d.max_points}
            for st in class_students(s, d.class_id) if st.id in assign]


def build_pdfs(did):
    import tempfile
    from pathlib import Path

    with db.session() as s:
        d, specs, assign = _load(s, did)
        cards = diagnostic_cards(s, did)
        cls = s.get(SchoolClass, d.class_id)
        title, dstr, class_name = d.title, fmt_date(d.date), cls.name
    with tempfile.TemporaryDirectory() as tmp:
        sheets, key = Path(tmp) / "v.pdf", Path(tmp) / "k.pdf"
        pdfgen.sheets_pdf(sheets, cards, dstr, title, class_name)
        pdfgen.key_pdf(key, specs, dstr, title)
        storage.put(pdf_key(did, "varaqlar"), sheets.read_bytes(), "application/pdf")
        storage.put(pdf_key(did, "kalit"), key.read_bytes(), "application/pdf")


def get_pdf(did, kind):
    if kind not in ("varaqlar", "kalit"):
        raise NotFound
    key = pdf_key(did, kind)
    if not storage.exists(key):
        build_pdfs(did)
    return storage.get(key)


def diagnostic_detail(did):
    with db.session() as s:
        d, specs, assign = _load(s, did)
        sts = class_students(s, d.class_id)
        resp = {r.student_id: r for r in s.scalars(select(Response).where(Response.diagnostic_id == did))}
        res = {r.student_id for r in s.scalars(select(Result).where(Result.diagnostic_id == did))}
        scans = s.scalars(select(Scan).where(Scan.diagnostic_id == did).order_by(Scan.id.desc())).all()
        by_no = {st.journal_no: st for st in sts}
        scan_counts, latest_scan = {}, {}
        for x in sorted(scans, key=lambda x: x.id):          # eskisidan yangisiga: oxirgisi ustun bo'ladi
            for n in (x.journal_nos or []):
                scan_counts[n] = scan_counts.get(n, 0) + 1
                latest_scan[n] = x.id
        rows = [{**student_brief(st), "level": assign.get(st.id),
                 "marks": resp[st.id].marks if st.id in resp else None,
                 "flags": resp[st.id].flags if st.id in resp else {},
                 "source": resp[st.id].source if st.id in resp else None,
                 "confidence": resp[st.id].confidence if st.id in resp else None,
                 "corrected": resp[st.id].corrected if st.id in resp else 0,
                 "graded": st.id in res} for st in sts]
        return {
            "id": d.id, "title": d.title, "template": d.template, "date": fmt_date(d.date), "status": d.status,
            "lesson_id": d.lesson_id, "class_id": d.class_id,
            "counts": {lvl: sum(1 for v in assign.values() if v == lvl) for lvl in LEVELS},
            "levels": [{"level": lvl, "name": LEVEL_NAMES[lvl], "spec": specs[lvl]} for lvl in LEVELS if lvl in specs],
            "question_steps": QUESTION_STEP, "errors": ERRORS,
            "rows": rows,
            "responses": sum(1 for r in rows if r["marks"]), "flagged": sum(1 for r in rows if r["flags"]),
            "graded": len(res),
            "pending": sum(1 for x in scans if x.status == "yuklandi"),
            "scans": [{"id": x.id, "status": x.status,
                       "url": storage_url(x.annotated_key or x.original_key), "original": storage_url(x.original_key),
                       "strips": x.strips_found,
                       "at": x.created_at.strftime("%H:%M") if x.created_at else "",
                       "students": [{"journal_no": n, "name": by_no[n].full_name, "code": by_no[n].code,
                                     "repeat": scan_counts.get(n, 0) > 1,
                                     "latest": latest_scan.get(n) == x.id}
                                    for n in (x.journal_nos or []) if n in by_no]}
                      for x in scans],
            "pdf": {"varaqlar": f"/api/diagnostics/{did}/pdf/varaqlar", "kalit": f"/api/diagnostics/{did}/pdf/kalit"},
            "sheets": -(-len(rows) // pdfgen.CARDS_PER_SHEET),
            "kind": d.kind, "max_points": d.max_points, "quality": scan_quality(did),
        }


def store_photo(did: int, data: bytes, filename: str = "surat.jpg") -> dict:
    """Suratni omborga yuklaydi va navbatga qo'yadi (hali o'qilmaydi)."""
    img = cv2.imdecode(np.frombuffer(data, np.uint8), cv2.IMREAD_COLOR)
    if img is None:
        raise ValueError(f"{filename}: rasmni o'qib bo'lmadi")
    with db.session() as s:
        _load(s, did)
    stamp = time.strftime("%Y%m%d-%H%M%S") + f"-{random.randint(100, 999)}"
    key = f"scans/{did}/{stamp}-asl.jpg"
    ok, buf = cv2.imencode(".jpg", img, [cv2.IMWRITE_JPEG_QUALITY, 85])
    storage.put(key, buf.tobytes(), "image/jpeg")
    with db.session() as s:
        s.add(Scan(diagnostic_id=did, original_key=key, annotated_key="", strips_found=0,
                   journal_nos=[], status="yuklandi"))
    return diagnostic_detail(did)


def scan_pending(did: int) -> dict:
    """Navbatdagi barcha suratlarni o'qiydi."""
    with db.session() as s:
        pending = [(x.id, x.original_key) for x in s.scalars(
            select(Scan).where(Scan.diagnostic_id == did, Scan.status == "yuklandi").order_by(Scan.id))]
    total = {"photos": 0, "found": 0, "matched": 0, "flagged": 0, "unknown": [], "errors": []}
    for scan_id, key in pending:
        try:
            res = process_photo(did, storage.get(key), key.rsplit("/", 1)[-1], scan_id=scan_id)
        except Exception as e:
            total["errors"].append(f"{key.rsplit('/', 1)[-1]}: {e}")
            continue
        if "error" in res:
            total["errors"].append(res["error"])
            continue
        total["photos"] += 1
        for k in ("found", "matched", "flagged"):
            total[k] += res[k]
        total["unknown"] += res["unknown"]
    return total


def process_photo(did: int, data: bytes, filename: str, scan_id: int | None = None) -> dict:
    img = cv2.imdecode(np.frombuffer(data, np.uint8), cv2.IMREAD_COLOR)
    if img is None:
        return {"error": f"{filename}: rasmni o'qib bo'lmadi"}
    reads, annotated = omr.scan_image(img)
    stamp = time.strftime("%Y%m%d-%H%M%S") + f"-{random.randint(100, 999)}"
    orig_key, ann_key = f"scans/{did}/{stamp}-asl.jpg", f"scans/{did}/{stamp}-belgilangan.jpg"
    ok, buf = cv2.imencode(".jpg", img, [cv2.IMWRITE_JPEG_QUALITY, 85])
    storage.put(orig_key, buf.tobytes(), "image/jpeg")
    h, w = annotated.shape[:2]
    k = min(1.0, 1800 / max(h, w))
    ok, buf = cv2.imencode(".jpg", cv2.resize(annotated, None, fx=k, fy=k), [cv2.IMWRITE_JPEG_QUALITY, 85])
    storage.put(ann_key, buf.tobytes(), "image/jpeg")

    matched, flagged, unknown = 0, 0, []
    with db.session() as s:
        d, specs, assign = _load(s, did)
        by_no = {st.journal_no: st for st in class_students(s, d.class_id)}
        for r in reads:
            st = by_no.get(r.journal_no)
            if not st or st.id not in assign:
                unknown.append(r.journal_no)
                continue
            row = s.get(Response, (did, st.id))
            if row:
                row.marks, row.flags, row.source = r.marks, r.flags, "skaner"
                row.auto_marks, row.confidence, row.corrected = dict(r.marks), r.confidence, 0
            else:
                s.add(Response(diagnostic_id=did, student_id=st.id, marks=r.marks, flags=r.flags, source="skaner",
                               auto_marks=dict(r.marks), confidence=r.confidence, corrected=0))
            matched += 1
            flagged += bool(r.flags)
        nos = sorted(r.journal_no for r in reads)
        row = s.get(Scan, scan_id) if scan_id else None
        if row:
            row.original_key, row.annotated_key = orig_key, ann_key
            row.strips_found, row.journal_nos, row.status = len(reads), nos, "o'qildi"
        else:
            s.add(Scan(diagnostic_id=did, original_key=orig_key, annotated_key=ann_key, strips_found=len(reads),
                       journal_nos=nos, status="o'qildi"))
        if d.status == "yaratildi":
            d.status = "skanerlandi"
    return {"found": len(reads), "matched": matched, "flagged": flagged, "unknown": unknown, "annotated": storage_url(ann_key)}


def delete_scan(did: int, scan_id: int) -> dict:
    """Yuklangan suratni o'chiradi (javoblar saqlanadi — ularni o'qituvchi qo'lda tuzatishi mumkin)."""
    with db.session() as s:
        row = s.get(Scan, scan_id)
        if not row or row.diagnostic_id != did:
            raise NotFound
        keys = [row.original_key, row.annotated_key]
        s.delete(row)
    for key in keys:
        try:
            storage.delete(key)
        except Exception:
            pass
    return diagnostic_detail(did)


def save_response(did: int, sid: int, marks: dict):
    marks = {k: (marks.get(k) or None) for k in ("q1", "q2", "q3", "q4", "q5", "q6", "q7")}
    if marks["q5"]:
        marks["q5"] = str(marks["q5"]).replace(".", ",").strip()
    with db.session() as s:
        _load(s, did)
        row = s.get(Response, (did, sid))
        if row:
            auto = row.auto_marks or {}
            row.corrected = sum(1 for k in marks if auto.get(k) != marks.get(k)) if auto else 0
            row.marks, row.flags, row.source = marks, {}, "qo'lda"
        else:
            s.add(Response(diagnostic_id=did, student_id=sid, marks=marks, flags={}, source="qo'lda"))


def demo_photo(did: int, count: int = 10) -> dict:
    """Printer bo'lmaganda: javobsiz o'quvchilar uchun sun'iy to'ldirilgan kartochkalar surati va uni skanerlash."""
    with db.session() as s:
        d, specs, assign = _load(s, did)
        done = set(s.scalars(select(Response.student_id).where(Response.diagnostic_id == did)))
        smap = skill_map(s, d.class_id)
        all_cards = {c["journal_no"]: c for c in diagnostic_cards(s, did)}
        pending = [(st.id, st.journal_no) for st in class_students(s, d.class_id) if st.id not in done][:count]
        cls = s.get(SchoolClass, d.class_id)
        title, dstr, class_name = d.title, fmt_date(d.date), cls.name
    if not pending:
        return {"error": "Barcha o'quvchilar javobi allaqachon kiritilgan."}
    rng = random.Random(did * 31 + len(done))
    cards, answers = [], {}
    for sid, jno in pending:
        sk = smap.get(sid, {})
        if sk.get("tayanch", 1) < 0.4:
            profile = "tayanch"
        elif avg(sk.values()) > 0.8:
            profile = "kuchli"
        else:
            profile = rng.choice(simulate.PROFILES)
        cards.append(all_cards[jno])
        answers[jno] = simulate.simulate_answers(specs[assign[sid]], profile, rng)
    img = simulate.make_photo(cards, answers, dstr, title, class_name, seed=rng.randint(0, 10 ** 6))
    ok, buf = cv2.imencode(".jpg", img, [cv2.IMWRITE_JPEG_QUALITY, 88])
    return process_photo(did, buf.tobytes(), "demo.jpg")


# ------------------------------------------------------------------ baholash
def grade_diagnostic(did: int, use_llm: bool = True) -> dict:
    t0 = time.time()
    with db.session() as s:
        d, specs, assign = _load(s, did)
        codes = {st.id: st.code for st in class_students(s, d.class_id)}
        smap = skill_map(s, d.class_id)
        gaps = attention_gaps(s, d.class_id)
        responses = [(r.student_id, r.marks) for r in s.scalars(select(Response).where(Response.diagnostic_id == did))]
        title, kind, max_points = d.title, d.kind, d.max_points

    graded, items = {}, []
    for sid, marks in responses:
        level = assign[sid]
        res = grading.grade(specs[level], marks)
        graded[sid] = (level, res)
        items.append({
            "code": codes[sid], "variant": f"{level} ({LEVEL_NAMES[level]})", "togri": f"{res['correct']}/{res['total']}",
            "bosqichlar": [{"bosqich": x["step"], "natija": "to'g'ri" if x["ok"] else "xato"} for x in res["steps"]],
            "asosiy_xato": res["primary_text"], "asosiy_xato_kodi": res["primary_error"], "izoh": res["root_cause_note"],
        })
    if not graded:
        return {"graded": 0, "feedback_source": "shablon", "summary_source": "shablon"}

    if use_llm:
        feedback, fb_source = llm.student_feedback(title, items)
    else:
        feedback, fb_source = {it["code"]: llm.fallback_feedback(it) for it in items}, "shablon"

    with db.session() as s:
        s.execute(delete(Result).where(Result.diagnostic_id == did))
        for sid, (level, res) in graded.items():
            fb = feedback[codes[sid]]
            s.add(Result(diagnostic_id=did, student_id=sid, correct=res["correct"], total=res["total"],
                         primary_error=res["primary_error"], details=res, feedback_student=fb["oquvchiga"],
                         feedback_parent=fb["ota_onaga"], feedback_teacher=fb["oqituvchiga"], feedback_source=fb_source))
            for skill, score in grading.update_skills(smap.get(sid, {}), res, level).items():
                row = s.get(SkillScore, (sid, skill))
                if row:
                    row.score = score
                else:
                    s.add(SkillScore(student_id=sid, skill=skill, score=score))
        s.get(Diagnostic, did).status = "baholandi"

    grades_from_diagnostic(did)
    stats = class_stats(graded, codes, gaps)
    stats["mavzu"] = title
    if use_llm:
        summary, s_source = llm.class_summary(stats)
    else:
        summary, s_source = {"xulosa": stats["fallback_xulosa"], "keyingi_dars": stats["fallback_keyingi"]}, "shablon"
    with db.session() as s:
        row = s.get(DiagSummary, did)
        payload = {"summary": summary, "stats": {k: v for k, v in stats.items() if not k.startswith("fallback")},
                   "seconds": round(time.time() - t0, 1)}
        if row:
            row.summary, row.source = payload, s_source
        else:
            s.add(DiagSummary(diagnostic_id=did, summary=payload, source=s_source))
    return {"graded": len(graded), "feedback_source": fb_source, "summary_source": s_source}


def rate_feedback(did: int, sid: int, rating: int | None = None, text: str | None = None):
    """Feedback sifati: 👍/👎 va o'qituvchi tahriri (AI matnining qanchasi o'zgarishsiz ketganini o'lchash uchun)."""
    with db.session() as s:
        row = s.get(Result, (did, sid))
        if not row:
            raise NotFound
        if rating is not None:
            row.feedback_rating = max(-1, min(int(rating), 1))
        if text is not None and text.strip() and text.strip() != (row.feedback_student or "").strip():
            row.feedback_student = text.strip()[:600]
            row.feedback_edited = True
    return diagnostic_results(did)


def scan_quality(did: int) -> dict:
    """Avtomatik o'qish ishonchliligi: nechta katak avtomatik o'qildi, nechtasi o'qituvchiga chiqdi va tuzatildi."""
    with db.session() as s:
        rows = s.scalars(select(Response).where(Response.diagnostic_id == did)).all()
        results = s.scalars(select(Result).where(Result.diagnostic_id == did)).all()
    cells = 7 * len(rows)
    unsure = sum(len([k for k in (r.flags or {}) if k != "chiziq"]) for r in rows)
    corrected = sum(r.corrected or 0 for r in rows)
    confs = [r.confidence for r in rows if r.confidence is not None]
    return {
        "students": len(rows), "cells": cells,
        "scanned": sum(1 for r in rows if r.auto_marks), "manual": sum(1 for r in rows if not r.auto_marks),
        "unsure": unsure, "corrected": corrected,
        "auto_pct": round(100 * (cells - unsure) / cells, 1) if cells else None,
        "accuracy_pct": round(100 * (cells - corrected) / cells, 1) if cells else None,
        "confidence_pct": round(100 * avg(confs)) if confs else None,
        "feedback_total": len(results),
        "feedback_edited": sum(1 for r in results if r.feedback_edited),
        "feedback_up": sum(1 for r in results if (r.feedback_rating or 0) > 0),
        "feedback_down": sum(1 for r in results if (r.feedback_rating or 0) < 0),
        "feedback_kept_pct": round(100 * (len(results) - sum(1 for r in results if r.feedback_edited)) / len(results))
        if results else None,
    }


def class_stats(graded, codes, gaps):
    n = len(graded)
    step_ok = {k: 0 for k in QUESTION_STEP}
    errors = {}
    for sid, (level, res) in graded.items():
        for x in res["steps"]:
            step_ok[x["key"]] += x["ok"]
        if res["primary_error"]:
            errors[res["primary_error"]] = errors.get(res["primary_error"], 0) + 1
    step_pct = {QUESTION_STEP[k]: round(100 * v / n) if n else 0 for k, v in step_ok.items()}
    N = gap_alert()
    neglected = sorted(((codes[sid], g) for sid, g in gaps.items() if g >= N), key=lambda x: -x[1])
    mastered = [codes[sid] for sid, (lvl, res) in graded.items() if res["correct"] == res["total"]]
    top_errors = sorted(errors.items(), key=lambda x: -x[1])
    avg_pct = round(100 * avg(res["correct"] / res["total"] for _, res in graded.values())) if n else 0

    fb_x = [f"{n} o'quvchi baholandi, o'rtacha natija {avg_pct}%."]
    if step_pct:
        weakest = min(step_pct.items(), key=lambda x: x[1])
        fb_x.append(f"Eng qiyin bosqich: “{weakest[0]}” — faqat {weakest[1]}% to'g'ri.")
    for err, cnt in top_errors[:2]:
        fb_x.append(f"{cnt} o'quvchida asosiy muammo: {ERRORS[err].lower()}.")
    fb_k = []
    if top_errors:
        fb_k.append(f"Takrorlash bosqichida: {ERROR_ACTIONS[top_errors[0][0]]}")
    if mastered:
        fb_k.append(f"Guruh ishida murabbiy roli: {', '.join(mastered[:5])}.")
    if neglected:
        fb_k.append("Alohida e'tibor: " + ", ".join(f"{c} ({g} dars)" for c, g in neglected[:5]) + ".")
    return {
        "mavzu": None, "baholangan_oquvchilar": n, "ortacha_foiz": avg_pct,
        "bosqichlar_boyicha_togri_foiz": step_pct,
        "asosiy_xatolar": {ERRORS[k]: v for k, v in top_errors},
        "asosiy_xato_kodlari": {k: v for k, v in top_errors},
        "toliq_ozlashtirganlar": mastered,
        "etiborsiz_oquvchilar": [{"kod": c, "dars": g} for c, g in neglected],
        "fallback_xulosa": fb_x, "fallback_keyingi": fb_k,
    }


def diagnostic_results(did):
    with db.session() as s:
        d, specs, assign = _load(s, did)
        sts = {st.id: st for st in class_students(s, d.class_id)}
        gaps = attention_gaps(s, d.class_id)
        results = s.scalars(select(Result).where(Result.diagnostic_id == did)).all()
        summary = s.get(DiagSummary, did)
        items = []
        for r in sorted(results, key=lambda r: sts[r.student_id].journal_no):
            st = sts[r.student_id]
            items.append({**student_brief(st), "level": assign.get(st.id), "gap": gaps.get(st.id, 0),
                          "correct": r.correct, "total": r.total, "primary_error": r.primary_error,
                          "primary_text": r.details["primary_text"], "root_cause_note": r.details.get("root_cause_note"),
                          "steps": [{"key": x["key"], "step": x["step"], "ok": x["ok"], "error": x["error"],
                                     "error_text": ERRORS.get(x["error"]) if x["error"] else None} for x in r.details["steps"]],
                          "feedback": {"student": r.feedback_student, "parent": r.feedback_parent, "teacher": r.feedback_teacher},
                          "feedback_source": r.feedback_source, "feedback_rating": r.feedback_rating,
                          "feedback_edited": r.feedback_edited})
        stats = summary.summary["stats"] if summary else None
        return {
            "id": d.id, "title": d.title, "date": fmt_date(d.date), "lesson_id": d.lesson_id, "class_id": d.class_id,
            "graded": len(items),
            "avg_pct": round(100 * avg(i["correct"] / i["total"] for i in items)) if items else None,
            "minutes_saved": round(len(items) * MINUTES_SAVED_PER_STUDENT),
            "quality": scan_quality(did),
            "kind": d.kind, "max_points": d.max_points,
            "summary": summary.summary["summary"] if summary else None,
            "summary_source": summary.source if summary else None,
            "seconds": summary.summary.get("seconds") if summary else None,
            # JSONB kalitlar tartibini saqlamaydi: bosqichlar savollar tartibida, xatolar soni bo'yicha
            "steps": [{"name": name, "pct": stats["bosqichlar_boyicha_togri_foiz"].get(name, 0)}
                      for name in QUESTION_STEP.values()] if stats else [],
            "errors": sorted(({"name": k, "count": v} for k, v in stats["asosiy_xatolar"].items()),
                             key=lambda e: -e["count"]) if stats else [],
            "students": items,
        }


# ------------------------------------------------------------------ dars ssenariysi
def active_group_membership(s, class_id, student_ids, max_age_days=14):
    """Oxirgi 2 hafta ichida tuzilgan guruhlar tarkibi (har darsda qayta bo'lmaslik uchun)."""
    row = s.scalars(select(ClassGroups).where(ClassGroups.class_id == class_id, ClassGroups.active.is_(True))
                    .order_by(ClassGroups.id.desc())).first()
    if not row or row.created_at < clock.now() - timedelta(days=max_age_days):
        return None
    alive = set(student_ids)
    groups = [[i for i in g if i in alive] for g in row.members]
    return [g for g in groups if g] or None


def create_lesson_plan(lesson_id: int | None = None, diagnostic_id: int | None = None, class_id: int | None = None,
                       use_llm: bool = True, regroup: bool = False) -> int:
    rng = random.Random(time.time_ns())
    N = gap_alert()
    prefs = method_prefs()
    with db.session() as s:
        if lesson_id:
            lesson = s.get(Lesson, lesson_id)
            if not lesson:
                raise NotFound
            cls = s.get(SchoolClass, lesson.class_id)
        else:
            cls = get_class(s, class_id)
            lesson = today_lesson(s, cls.id)
        sts = class_students(s, cls.id)
        if diagnostic_id:
            d = s.get(Diagnostic, diagnostic_id)
        else:
            d = s.scalars(select(Diagnostic).where(Diagnostic.class_id == cls.id, Diagnostic.status == "baholandi",
                                                   Diagnostic.date <= lesson.date)
                          .order_by(Diagnostic.date.desc(), Diagnostic.id.desc())).first()
        results, focus = {}, {}
        topic = lesson.topic or (d.title if d else "Matematika")
        from . import curriculum_plan
        topic_row = s.get(CurriculumTopic, lesson.curriculum_topic_id) if lesson.curriculum_topic_id else None
        refs = curriculum_plan.topic_refs(topic_row)
        template = lesson_template(s, lesson)
        if d:
            results = {r.student_id: r.details for r in s.scalars(select(Result).where(Result.diagnostic_id == d.id))}
            summ = s.get(DiagSummary, d.id)
            if summ:
                stats = summ.summary["stats"]
                weakest = min(stats["bosqichlar_boyicha_togri_foiz"].items(), key=lambda x: x[1])
                codes = stats.get("asosiy_xato_kodlari") or {}
                top = max(codes, key=codes.get) if codes else None
                focus = {"eng_zaif_bosqich": weakest[0], "togri_foiz": weakest[1],
                         "asosiy_xato": ERRORS.get(top) if top else None,
                         "tavsiya": ERROR_ACTIONS.get(top) if top else None,
                         "ortacha_foiz": stats.get("ortacha_foiz"), "diagnostika": d.title}
        # o'tgan darslardagi uy vazifasi xatolari ham keyingi ssenariyga kiradi
        hw_rows = s.scalars(select(Homework).join(Lesson, Lesson.id == Homework.lesson_id)
                            .where(Lesson.class_id == cls.id, Lesson.date < lesson.date)
                            .order_by(Homework.id.desc()).limit(40)).all()
        hw_errors = {}
        for row in hw_rows:
            for t in (row.tasks or []):
                if not t.get("togri") and t.get("xato"):
                    hw_errors[t["xato"]] = hw_errors.get(t["xato"], 0) + 1
        if hw_errors:
            top_hw = max(hw_errors, key=hw_errors.get)
            focus["uy_vazifasi_xatosi"] = f"{top_hw} ({hw_errors[top_hw]} ta mashqda)"
        prev = s.scalars(select(Lesson).where(Lesson.class_id == cls.id, Lesson.date < lesson.date, Lesson.quick_check.is_not(None))
                         .order_by(Lesson.date.desc())).first()
        struggled = set((prev.quick_check or {}).get("struggled", [])) if prev else set()
        rows = scenario.priority([student_brief(x) for x in sts], skill_map(s, cls.id), attention_gaps(s, cls.id),
                                 results, N, struggled)
        fmt = scenario.lesson_format(lesson.diagnostic_day, lesson.group_work)
        membership = None if regroup or not fmt["guruh_ishi"] else active_group_membership(s, cls.id, [x.id for x in sts])
        class_id, n_students, did, lid = cls.id, len(sts), (d.id if d else None), lesson.id

    targets = scenario.pick_targets(rows)
    problem = scenario.group_problem(rng, template) if fmt["guruh_ishi"] else None
    groups, new_membership = [], None
    if fmt["guruh_ishi"]:
        by_id = {r["id"]: r for r in rows}
        if membership:
            groups_rows = [[by_id[i] for i in g if i in by_id] for g in membership]
            placed = {i for g in membership for i in g}
            for r in rows:  # yangi qo'shilgan o'quvchilar eng kichik guruhga
                if r["id"] not in placed:
                    min(groups_rows, key=len).append(r)
        else:
            groups_rows = scenario.form_groups(rows)
            new_membership = [[r["id"] for r in g] for g in groups_rows]
        groups = scenario.describe_groups(groups_rows)
    context = scenario.build_context(topic, n_students, focus, targets, problem, rng, prefs, fmt, refs)
    llm_plan = llm.compose_lesson(context) if use_llm else None
    code_to_name = {r["code"]: r["name"] for r in rows}
    plan = scenario.finalize(llm_plan, topic, focus, targets, groups, problem, code_to_name, fmt, refs, template)
    if new_membership:
        plan["groups_formed"] = clock.today().strftime("%d.%m.%Y")
    with db.session() as s:
        if new_membership:
            for old in s.scalars(select(ClassGroups).where(ClassGroups.class_id == class_id, ClassGroups.active.is_(True))):
                old.active = False
            s.add(ClassGroups(class_id=class_id, members=new_membership, active=True))
        lp = LessonPlan(class_id=class_id, diagnostic_id=did, lesson_id=lid, topic=topic, plan=plan,
                        source="gpt" if llm_plan else "shablon")
        s.add(lp)
        s.flush()
        return lp.id


def lesson_plans(class_id=None):
    with db.session() as s:
        q = select(LessonPlan)
        if class_id:
            q = q.where(LessonPlan.class_id == class_id)
        return [{"id": p.id, "topic": p.topic, "source": p.source, "diagnostic_id": p.diagnostic_id, "lesson_id": p.lesson_id,
                 "class_id": p.class_id, "created_at": p.created_at.strftime("%d.%m.%Y %H:%M")}
                for p in s.scalars(q.order_by(LessonPlan.id.desc()))]


def lesson_plan(pid):
    prefs = method_prefs()
    with db.session() as s:
        p = s.get(LessonPlan, pid)
        if not p:
            raise NotFound
        lesson = s.get(Lesson, p.lesson_id) if p.lesson_id else today_lesson(s, p.class_id)
        marked = {st.code for st in s.scalars(select(Student).join(Attention, Attention.student_id == Student.id)
                                              .where(Attention.lesson_id == lesson.id, Attention.kind == "ishladim"))}
        plan = dict(p.plan)
        fmt = plan.get("format") or scenario.lesson_format(True, True)
        plan["format"] = fmt
        plan["stages"] = [{**st, "candidates": [{"id": m["id"], "name": m["name"], "short": m["short"], "custom": m.get("custom", False)}
                                                for m in scenario.stage_candidates(st["key"], fmt, prefs)]}
                          for st in plan.get("stages", [])]
        return {"id": p.id, "source": p.source, "diagnostic_id": p.diagnostic_id, "lesson_id": lesson.id, "class_id": p.class_id,
                "created_at": p.created_at.strftime("%d.%m.%Y %H:%M"), "status": plan.pop("status", "qoralama"),
                "attended_today": sorted(marked), **plan}


# ------------------------------------------------------------------ baholar jurnali (FR-23)
GRADE_KINDS = {"formativ": 10, "bsb": 50, "chsb": 40}


def grades_view(class_id=None, limit=20):
    """Sinfning baholar jurnali: oxirgi darslar × o'quvchilar."""
    from . import curriculum_plan

    with db.session() as s:
        cls = get_class(s, class_id)
        lessons = s.scalars(select(Lesson).where(Lesson.class_id == cls.id, Lesson.conducted_at.is_not(None))
                            .order_by(Lesson.date.desc()).limit(limit)).all()
        lessons = list(reversed(lessons))
        ids = [x.id for x in lessons]
        rows = s.scalars(select(Grade).where(Grade.lesson_id.in_(ids))).all() if ids else []
        by_student = {}
        for g in rows:
            by_student.setdefault(g.student_id, {})[g.lesson_id] = g
        topics = {t.id: t for t in s.scalars(select(CurriculumTopic).where(CurriculumTopic.class_id == cls.id))}
        cols = []
        for x in lessons:
            topic = topics.get(x.curriculum_topic_id)
            kind = "formativ"
            if topic and topic.kind in ("bsb", "chsb"):
                kind = topic.kind
            cols.append({"lesson_id": x.id, "date": fmt_date(x.date), "short": x.date.strftime("%d.%m"),
                         "topic": x.topic, "kind": kind, "max": topic.points if topic and topic.points else GRADE_KINDS[kind],
                         "lesson_no": topic.lesson_no if topic else None,
                         "diagnostic_day": x.diagnostic_day})
        students = []
        for st in class_students(s, cls.id):
            marks = by_student.get(st.id, {})
            vals = [round(100 * g.points / max(1, g.max_points)) for g in marks.values()]
            students.append({**student_brief(st),
                             "grades": {str(lid): {"points": g.points, "max": g.max_points, "kind": g.kind, "source": g.source}
                                        for lid, g in marks.items()},
                             "avg": round(sum(vals) / len(vals)) if vals else None})
        return {"class": cls.name, "class_id": cls.id, "lessons": cols, "students": students,
                "quarter": curriculum_plan.aha_plan.quarter_of(clock.today())[0]}


def save_grades(lesson_id: int, marks: dict, kind: str = "formativ", source: str = "qo'lda"):
    """marks: {student_id: ball yoki None}. None — bahoni o'chiradi."""
    if kind not in GRADE_KINDS:
        raise ValueError("Baho turi noto'g'ri")
    with db.session() as s:
        lesson = s.get(Lesson, lesson_id)
        if not lesson:
            raise NotFound
        topic = s.get(CurriculumTopic, lesson.curriculum_topic_id) if lesson.curriculum_topic_id else None
        if topic and topic.kind in ("bsb", "chsb"):
            kind = topic.kind
        top = (topic.points if topic and topic.points else GRADE_KINDS[kind])
        ids = {x.id for x in class_students(s, lesson.class_id)}
        for sid, value in marks.items():
            sid = int(sid)
            if sid not in ids:
                continue
            row = s.get(Grade, (lesson_id, sid))
            if value is None or value == "":
                if row:
                    s.delete(row)
                continue
            points = max(0.0, min(float(value), float(top)))
            if row:
                row.points, row.max_points, row.kind, row.source = points, top, kind, source
            else:
                s.add(Grade(lesson_id=lesson_id, student_id=sid, points=points, max_points=top, kind=kind, source=source))
        cid = lesson.class_id
    return grades_view(cid)


def grades_from_diagnostic(did: int):
    """Baholangan diagnostikadan ballarni jurnalga yozadi: formativ 0–10, summativ esa BSB/ChSB maksimaliga nisbatan."""
    with db.session() as s:
        d = s.get(Diagnostic, did)
        if not d or not d.lesson_id:
            return 0
        results = s.scalars(select(Result).where(Result.diagnostic_id == did)).all()
        kind, top, lid = d.kind, d.max_points, d.lesson_id
        if kind == "formativ":
            marks = {r.student_id: grading.formative_points(r.correct, r.total)
                     for r in results}
        else:
            marks = {r.student_id: round(top * grading.total_ratio(r.correct, r.total), 1)
                     for r in results}
    if marks:
        save_grades(lid, marks, kind=kind, source="diagnostika")
    return len(marks)


# ------------------------------------------------------------------ uy vazifasi (mashq daftari surati)
def homework_reference(s, lesson):
    topic = s.get(CurriculumTopic, lesson.curriculum_topic_id) if lesson.curriculum_topic_id else None
    return (topic.workbook if topic else None), (topic.title if topic else lesson.topic)


def homework_view(lesson_id: int):
    with db.session() as s:
        lesson = s.get(Lesson, lesson_id)
        if not lesson:
            raise NotFound
        ref, topic = homework_reference(s, lesson)
        rows = {h.student_id: h for h in s.scalars(select(Homework).where(Homework.lesson_id == lesson_id))}
        students = [{**student_brief(st),
                     "homework": ({"id": rows[st.id].id, "status": rows[st.id].status,
                                   "correct": rows[st.id].correct, "total": rows[st.id].total,
                                   "tasks": rows[st.id].tasks, "comment": rows[st.id].comment,
                                   "source": rows[st.id].source, "confirmed": rows[st.id].confirmed,
                                   "image": storage_url(rows[st.id].image_key) if rows[st.id].image_key else None}
                                  if st.id in rows else None)}
                    for st in class_students(s, lesson.class_id)]
        checked = [x for x in students if x["homework"] and x["homework"]["status"] == "tayyor"]
        errors = {}
        for x in checked:
            for t in x["homework"]["tasks"]:
                if not t.get("togri") and t.get("xato"):
                    errors[t["xato"]] = errors.get(t["xato"], 0) + 1
        pending = sum(1 for x in students if x["homework"] and x["homework"]["status"] == "navbatda")
        return {"lesson_id": lesson_id, "class_id": lesson.class_id, "date": fmt_date(lesson.date), "pending": pending,
                "topic": topic, "reference": ref, "students": students,
                "checked": len(checked),
                "avg_pct": round(100 * avg(x["homework"]["correct"] / max(1, x["homework"]["total"]) for x in checked))
                if checked else None,
                "minutes_saved": round(len(checked) * MINUTES_SAVED_PER_HOMEWORK),
                "top_errors": sorted(({"name": k, "count": v} for k, v in errors.items()), key=lambda e: -e["count"])[:4]}


PAGE_MESSAGE = {
    "kartochka": "Bu — diagnostika kartochkasi surati, uy vazifasi emas. Mashq daftari sahifasini suratga oling.",
    "darslik": "Suratda bosma kitob sahifasi ko'rinyapti — o'quvchining yozgan ishi yo'q.",
    "boshqa": "Suratda daftar sahifasi topilmadi. Yaqinroqdan, yorug'roq joyda qayta suratga oling.",
}


def check_homework(lesson_id: int, student_id: int, data: bytes, filename: str = "uy.jpg") -> dict:
    """Suratni saqlaydi va tahlilni fonga qo'yadi — o'qituvchi keyingi o'quvchiga o'tishi mumkin."""
    img = cv2.imdecode(np.frombuffer(data, np.uint8), cv2.IMREAD_COLOR)
    if img is None:
        raise ValueError(f"{filename}: rasmni o'qib bo'lmadi")
    h, w = img.shape[:2]
    k = min(1.0, 1700 / max(h, w))
    if k < 1.0:
        img = cv2.resize(img, None, fx=k, fy=k, interpolation=cv2.INTER_AREA)
    ok, buf = cv2.imencode(".jpg", img, [cv2.IMWRITE_JPEG_QUALITY, 85])
    jpg = buf.tobytes()
    key = f"homework/{lesson_id}/{student_id}.jpg"
    storage.put(key, jpg, "image/jpeg")
    with db.session() as s:
        lesson = s.get(Lesson, lesson_id)
        if not lesson:
            raise NotFound
        ref, _topic = homework_reference(s, lesson)
        row = s.scalars(select(Homework).where(Homework.lesson_id == lesson_id, Homework.student_id == student_id)).first()
        payload = dict(image_key=key, reference=ref, tasks=[], correct=0, total=0,
                       comment="Navbatda — AI tekshirmoqda…", source="ai", confirmed=False, status="navbatda")
        if row:
            for k2, v in payload.items():
                setattr(row, k2, v)
        else:
            s.add(Homework(lesson_id=lesson_id, student_id=student_id, **payload))
    _BG.submit(analyze_homework, lesson_id, student_id, jpg)
    return homework_view(lesson_id)


DEMO_HOMEWORK_OK = ["24 + 6 * 3 = 42", "(120 - 45) : 5 = 15", "36 : 4 + 18 = 27", "7 * 8 - 14 = 42"]
DEMO_HOMEWORK_BAD = ["24 + 6 * 3 = 90", "(120 - 45) : 5 = 105", "36 : 4 + 18 = 25", "150 - 60 : 6 = 15"]


def demo_homework(lesson_id: int, student_id: int | None = None) -> dict:
    """Daftar bo'lmaganda: sun'iy mashq daftari sahifasi va uni haqiqiy AI tekshiruvidan o'tkazish.

    Surat chindan ham generatsiya qilinadi va oddiy uy vazifasi kabi yo'lga qo'yiladi —
    demoda mentor AI ning suratni o'qishini to'liq ko'radi.
    """
    from . import simulate

    with db.session() as s:
        lesson = s.get(Lesson, lesson_id)
        if not lesson:
            raise NotFound
        students = class_students(s, lesson.class_id)
        have = {h.student_id for h in s.scalars(select(Homework).where(Homework.lesson_id == lesson_id))}
        if student_id:
            st = next((x for x in students if x.id == student_id), None)
            if st is None:
                raise ValueError("Bu o'quvchi shu sinfda emas")
        else:
            st = next((x for x in students if x.id not in have), None)
            if st is None:
                return {"error": "Sinfdagi barcha o'quvchilarning vazifasi allaqachon yuklangan."}
        sid, jno = st.id, st.journal_no
        skills = skill_map(s, lesson.class_id).get(sid, {})

    rng = random.Random(lesson_id * 97 + sid)
    strong = avg(skills.values()) > 0.75 if skills else rng.random() < 0.5
    wrong = 0 if strong else rng.randint(1, 2)          # kuchli o'quvchida xato yo'q, boshqasida 1–2 ta
    picks = rng.sample(range(len(DEMO_HOMEWORK_OK)), 3)
    tasks = []
    for k, i in enumerate(picks):
        bad = k < wrong
        tasks.append({"nom": str(112 + i * 2), "yozuv": (DEMO_HOMEWORK_BAD if bad else DEMO_HOMEWORK_OK)[i]})
    jpg = simulate.make_homework_page(tasks, seed=lesson_id * 10 + sid)
    out = check_homework(lesson_id, sid, jpg, f"demo-daftar-{jno}.jpg")
    return {**out, "demo_student_id": sid}


def analyze_homework(lesson_id: int, student_id: int, jpg: bytes):
    """Fon ishi: surat turini tekshiradi, so'ng AI mashqlarni baholaydi."""
    try:
        with db.session() as s:
            lesson = s.get(Lesson, lesson_id)
            ref, topic = homework_reference(s, lesson) if lesson else (None, None)
        # 1) tez tekshiruv: bizning javob blokimiz markerlari bormi? unda bu kartochka, daftar emas
        img = cv2.imdecode(np.frombuffer(jpg, np.uint8), cv2.IMREAD_COLOR)
        reads, _ann = omr.scan_image(img) if img is not None else ([], None)
        if reads:
            _finish_homework(lesson_id, student_id, None, "xato", PAGE_MESSAGE["kartochka"])
            return
        context = {"mavzu": topic, "manba": ref, "sinf": 5}
        got = llm.check_homework(jpg, context)
        # chegaradagi suratlarda tasnif bir marta adashishi mumkin — bir marta qayta so'raymiz
        if got is not None and not got.get("masalalar") and str(got.get("sahifa", "")) not in ("kartochka", "darslik"):
            got = llm.check_homework(jpg, context) or got
        if got is None:
            _finish_homework(lesson_id, student_id, None, "xato", "AI tekshira olmadi — qo'lda kiriting.")
            return
        if not got.get("masalalar"):
            page = str(got.get("sahifa", "boshqa"))
            _finish_homework(lesson_id, student_id, None, "xato",
                             PAGE_MESSAGE.get(page, got.get("izoh") or PAGE_MESSAGE["boshqa"]))
            return
        note = got["izoh"]
        if got.get("mavzuga_mos") is False:
            note = f"Diqqat: mashqlar bugungi mavzuga mos kelmasligi mumkin. {note}"
        _finish_homework(lesson_id, student_id, got, "tayyor", note)
    except Exception as e:                       # fon ishi ilovani yiqitmasligi kerak
        _finish_homework(lesson_id, student_id, None, "xato", f"Tahlilda xato: {type(e).__name__}")


def _finish_homework(lesson_id: int, student_id: int, got: dict | None, status: str, comment: str):
    with db.session() as s:
        row = s.scalars(select(Homework).where(Homework.lesson_id == lesson_id,
                                               Homework.student_id == student_id)).first()
        if not row:
            return
        row.tasks = got["masalalar"] if got else []
        row.correct = got["correct"] if got else 0
        row.total = got["total"] if got else 0
        row.comment = comment[:300]
        row.source = "ai" if got else "yo'q"
        row.status = status


def confirm_homework(lesson_id: int, student_id: int, tasks: list | None = None, confirmed: bool = True):
    with db.session() as s:
        row = s.scalars(select(Homework).where(Homework.lesson_id == lesson_id, Homework.student_id == student_id)).first()
        if not row:
            raise NotFound
        if tasks is not None:
            clean = [{"nom": str(t.get("nom", ""))[:10], "togri": bool(t.get("togri")),
                      "xato": (str(t.get("xato"))[:60] if t.get("xato") else None), "izoh": str(t.get("izoh", ""))[:160]}
                     for t in tasks if str(t.get("nom", "")).strip()]
            row.tasks, row.total = clean, len(clean)
            row.correct = sum(1 for t in clean if t["togri"])
            row.source = "o'qituvchi"
            row.status = "tayyor"
        row.confirmed = confirmed
    return homework_view(lesson_id)


# ------------------------------------------------------------------ ta'sir: tejalgan vaqt va qog'oz
def impact_stats(class_id=None):
    """Chorak boshidan beri AI nima qildi: nechta ish avtomatik tekshirildi, qancha vaqt va qog'oz tejaldi."""
    from . import conveyor
    from .reports import quarter_start

    q0 = quarter_start(clock.today())
    with db.session() as s:
        classes = [get_class(s, class_id)] if class_id else s.scalars(select(SchoolClass).order_by(SchoolClass.id)).all()
        ids = [c.id for c in classes]
        paper = conveyor.paper_stats(s, classes)
        diag_ids = [d.id for d in s.scalars(select(Diagnostic).where(Diagnostic.class_id.in_(ids), Diagnostic.date >= q0))]
        results = s.scalars(select(Result).where(Result.diagnostic_id.in_(diag_ids))).all() if diag_ids else []
        responses = s.scalars(select(Response).where(Response.diagnostic_id.in_(diag_ids))).all() if diag_ids else []
        hw = s.scalars(select(Homework).join(Lesson, Lesson.id == Homework.lesson_id)
                       .where(Lesson.class_id.in_(ids), Lesson.date >= q0)).all()
    cells = 7 * len(responses)
    corrected = sum(r.corrected or 0 for r in responses)
    minutes = round(len(results) * MINUTES_SAVED_PER_STUDENT + len(hw) * MINUTES_SAVED_PER_HOMEWORK)
    return {
        "since": paper["since"],
        "graded_works": len(results), "homework_checked": len(hw),
        "cells_read": cells, "cells_corrected": corrected,
        "accuracy_pct": round(100 * (cells - corrected) / cells, 1) if cells else None,
        "minutes_saved": minutes, "hours_saved": round(minutes / 60, 1),
        "feedback_written": sum(1 for r in results if r.feedback_student),
        "feedback_kept_pct": round(100 * (len(results) - sum(1 for r in results if r.feedback_edited)) / len(results))
        if results else None,
        "sheets_used": paper["sheets_used"], "sheets_if_every_lesson": paper["sheets_if_every_lesson"],
        "sheets_saved": paper["sheets_saved"], "paper_lessons": paper["paper_lessons"], "quick_lessons": paper["quick_lessons"],
    }


# ------------------------------------------------------------------ e'tibor jurnali, metodlar, AI jurnali
def attention_view(class_id=None, lesson_id=None):
    N = gap_alert()
    with db.session() as s:
        if lesson_id:
            lesson = s.get(Lesson, lesson_id)
            if not lesson:
                raise NotFound
            cls = s.get(SchoolClass, lesson.class_id)
        else:
            cls = get_class(s, class_id)
            lesson = today_lesson(s, cls.id)
        today = set(s.scalars(select(Attention.student_id).where(Attention.lesson_id == lesson.id, Attention.kind == "ishladim")))
        gaps = attention_gaps(s, cls.id)
        return {"class": cls.name, "class_id": cls.id, "gap_alert": N,
                "lesson": {"id": lesson.id, "date": fmt_date(lesson.date), "topic": lesson.topic},
                "students": [{**student_brief(st), "today": st.id in today, "gap": gaps[st.id]} for st in class_students(s, cls.id)]}


def save_attention(student_ids, class_id=None, lesson_id=None):
    with db.session() as s:
        if lesson_id:
            lesson = s.get(Lesson, lesson_id)
            if not lesson:
                raise NotFound
        else:
            lesson = today_lesson(s, get_class(s, class_id).id)
        s.execute(delete(Attention).where(Attention.lesson_id == lesson.id, Attention.kind == "ishladim"))
        enrolled = set(s.scalars(select(Attention.student_id).where(Attention.lesson_id == lesson.id)))
        for sid in set(student_ids) - enrolled:
            s.add(Attention(student_id=sid, lesson_id=lesson.id))


def methods_catalog():
    chosen = {}
    for (stage, mid), score in method_prefs().items():
        chosen[mid] = chosen.get(mid, 0) + max(score, 0)
    return {"stages": STAGES,
            "methods": [{**m, "custom": m.get("custom", False), "chosen": chosen.get(m["id"], 0)} for m in methods_mod.METHODS]}


def ai_log(limit=60):
    with db.session() as s:
        rows = s.scalars(select(LlmCall).order_by(LlmCall.id.desc()).limit(limit)).all()
        totals = s.execute(select(func.count(LlmCall.id), func.sum(LlmCall.prompt_tokens), func.sum(LlmCall.completion_tokens),
                                  func.avg(LlmCall.seconds))).one()
        return {
            "totals": {"calls": totals[0] or 0, "prompt_tokens": int(totals[1] or 0), "completion_tokens": int(totals[2] or 0),
                       "avg_seconds": round(float(totals[3] or 0), 1)},
            "calls": [{"id": r.id, "at": r.created_at.strftime("%d.%m %H:%M:%S"), "purpose": r.purpose, "ok": r.ok,
                       "seconds": r.seconds, "prompt_tokens": r.prompt_tokens, "completion_tokens": r.completion_tokens,
                       "error": r.error} for r in rows],
        }


# ------------------------------------------------------------------ eksport
def export_xlsx(did) -> bytes:
    res = diagnostic_results(did)
    wb = Workbook()
    ws = wb.active
    ws.title = "Natijalar"
    ws.append([f"{res['title']} · {res['date']} · formativ baholash uchun"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(["Jurnal №", "Kod", "F.I.Sh", "Variant", "To'g'ri", "Foiz", "Formativ ball (0–10)",
               "Asosiy muammo", "O'qituvchiga tavsiya", "Ota-onaga xabar"])
    for cell in ws[2]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3A5F")
    for it in res["students"]:
        ws.append([it["journal_no"], it["code"], it["name"], it["level"], f"{it['correct']}/{it['total']}",
                   round(100 * it["correct"] / it["total"]), grading.formative_points(it["correct"], it["total"]),
                   it["primary_text"], it["feedback"]["teacher"], it["feedback"]["parent"]])
    for i, w in enumerate([9, 9, 26, 9, 9, 7, 12, 40, 50, 60]):
        ws.column_dimensions[chr(65 + i)].width = w
    for row in ws.iter_rows(min_row=3):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------ ssenariyni tahrirlash (FR-07)
def method_prefs() -> dict:
    with db.session() as s:
        return {(r.stage, r.method_id): r.score for r in s.scalars(select(MethodPreference))}


def _bump_pref(s, stage, method_id, delta):
    row = s.get(MethodPreference, (stage, method_id))
    if row:
        row.score += delta
    else:
        s.add(MethodPreference(stage=stage, method_id=method_id, score=delta))


def update_lesson_plan(pid: int, stages_in: list, status: str | None):
    with db.session() as s:
        p = s.get(LessonPlan, pid)
        if not p:
            raise NotFound
        plan = dict(p.plan)
        by_key = {x.get("key"): x for x in stages_in if isinstance(x, dict)}
        new_stages, start = [], 0
        for st in plan["stages"]:
            st = dict(st)
            edit = by_key.get(st["key"])
            if edit:
                mid = edit.get("method_id")
                if mid and mid != st["method"]["id"]:
                    if mid not in {m["id"] for m in scenario.stage_candidates(st["key"], plan.get("format") or scenario.lesson_format(True, True))}:
                        raise ValueError("Bu metod ushbu bosqichga mos emas")
                    m = methods_mod.METHOD_BY_ID[mid]
                    _bump_pref(s, st["key"], st["method"]["id"], -1)
                    _bump_pref(s, st["key"], mid, +1)
                    st["method"] = {"id": m["id"], "name": m["name"], "short": m["short"], "form": m["form"], "steps": m["steps"]}
                    st["why"] = "O'qituvchi tanlovi: " + m["short"]
                if edit.get("minutes") is not None:
                    st["minutes"] = max(1, min(30, int(edit["minutes"])))
                removed = set(edit.get("removed_codes") or [])
                if removed:
                    st["targeted"] = [t for t in st["targeted"] if t["code"] not in removed]
            st["start"] = start
            start += st["minutes"]
            new_stages.append(st)
        plan["stages"] = new_stages
        plan["total_minutes"] = start
        plan["edited"] = True
        if status in ("qoralama", "tasdiqlangan"):
            plan["status"] = status
        p.plan = plan
    return lesson_plan(pid)


def toggle_attention(student_id: int, on: bool, lesson_id: int | None = None):
    """Ssenariy ichidan "bu o'quvchi bilan ishladim" belgisi (FR-18)."""
    with db.session() as s:
        st = s.get(Student, student_id)
        if not st:
            raise NotFound
        lesson = s.get(Lesson, lesson_id) if lesson_id else today_lesson(s, st.class_id)
        row = s.scalars(select(Attention).where(Attention.lesson_id == lesson.id, Attention.student_id == student_id)).first()
        if on and not row:
            s.add(Attention(student_id=student_id, lesson_id=lesson.id, kind="ishladim"))
        elif on and row:
            row.kind = "ishladim"
        elif not on and row and row.kind == "ishladim":
            s.delete(row)


def student_id_by_code(code: str) -> int:
    with db.session() as s:
        st = s.scalars(select(Student).where(Student.code == code)).first()
        if not st:
            raise NotFound
        return st.id


# ------------------------------------------------------------------ skaner: bir tegishda tasdiqlash (FR-15)
def confirm_response(did: int, sid: int):
    with db.session() as s:
        row = s.get(Response, (did, sid))
        if not row:
            raise NotFound
        row.flags = {}


# ------------------------------------------------------------------ o'quvchilar (FR-02)
def add_students(names: list, class_id=None) -> dict:
    names = [" ".join(str(n).split())[:120] for n in names]
    names = [n for n in names if any(ch.isalpha() for ch in n)]
    if not names:
        raise ValueError("Ism topilmadi")
    added = []
    with db.session() as s:
        cls = get_class(s, class_id)
        lesson = today_lesson(s, cls.id)
        last = s.scalar(select(func.max(Student.journal_no)).where(Student.class_id == cls.id)) or 0
        prefix = cls.name.replace("-", "")
        for i, name in enumerate(names, start=1):
            no = last + i
            st = Student(class_id=cls.id, journal_no=no, full_name=name, code=f"{prefix}-{no:02d}",
                         parent_token=secrets.token_urlsafe(9))
            s.add(st)
            s.flush()
            # e'tibor hisobi o'quvchi qo'shilgan kundan boshlanadi
            s.add(Attention(student_id=st.id, lesson_id=lesson.id, kind="qo'shildi"))
            added.append({"id": st.id, "code": st.code, "name": name})
    return {"added": added}


def import_students(data: bytes, filename: str, class_id=None) -> dict:
    rows = []
    if filename.lower().endswith(".xlsx"):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True)
        for r in wb.active.iter_rows(values_only=True):
            cells = [str(c).strip() for c in r if c is not None and any(ch.isalpha() for ch in str(c))]
            if cells:
                rows.append(max(cells, key=len))
    else:
        for line in data.decode("utf-8", errors="ignore").splitlines():
            parts = [x.strip() for x in line.replace(";", ",").split(",") if any(ch.isalpha() for ch in x)]
            if parts:
                rows.append(max(parts, key=len))
    if rows and any(w in rows[0].lower() for w in ("ism", "f.i", "familiya", "name")):
        rows = rows[1:]
    return add_students(rows, class_id)


# ------------------------------------------------------------------ o'qituvchi metodi (FR-10)
def _card_from_row(row: CustomMethod) -> dict:
    return {**row.card, "id": f"u{row.id}", "custom": True}


def load_custom_methods():
    with db.session() as s:
        for row in s.scalars(select(CustomMethod)):
            methods_mod.register(_card_from_row(row))


def validate_card(card: dict) -> dict:
    keys = {st["key"] for st in STAGES}
    stages = [x for x in card.get("stages", []) if x in keys] if isinstance(card.get("stages"), list) else []
    steps = [str(x).strip()[:200] for x in card.get("steps", []) if str(x).strip()] if isinstance(card.get("steps"), list) else []
    name = str(card.get("name") or "").strip()[:60]
    if not name or not stages:
        raise ValueError("Metod nomi va kamida bitta bosqich kerak")
    form = str(card.get("form") or "Sinf").strip()
    return {"name": name, "short": str(card.get("short") or name).strip()[:160], "stages": stages[:3],
            "form": form if form in ("Sinf", "Guruh", "Juftlik", "Yakka") else "Sinf",
            "minutes": str(card.get("minutes") or "3–5").strip()[:10], "steps": steps[:6] or [name]}


def method_draft(text: str) -> dict:
    text = (text or "").strip()
    if len(text) < 15:
        raise ValueError("Metodni kamida bir-ikki gap bilan tasvirlang")
    card, source = llm.method_card(text, STAGES), "gpt"
    try:
        if card:
            return {"card": validate_card(card), "source": source}
    except ValueError:
        pass
    low = text.lower()
    sentences = [x.strip() for x in text.replace("\n", ". ").split(".") if x.strip()]
    stage_words = {"takrorlash": ["takror"], "yangi_mavzu": ["yangi", "tushuntir"], "mustahkamlash": ["mustahkam", "guruh"],
                   "refleksiya": ["refleks", "xulosa"], "motivatsiya": ["qiziq", "boshla"]}
    guess = [k for k, words in stage_words.items() if any(w in low for w in words)] or ["mustahkamlash"]
    card = {"name": sentences[0][:40], "short": text[:120], "stages": guess[:2],
            "form": "Guruh" if "guruh" in low else "Juftlik" if "juft" in low else "Sinf",
            "minutes": "3–5", "steps": sentences[:4] or [text[:120]]}
    return {"card": validate_card(card), "source": "shablon"}


def save_method(card: dict, source_text: str | None) -> dict:
    card = validate_card(card)
    with db.session() as s:
        row = CustomMethod(card=card, source_text=(source_text or "")[:2000])
        s.add(row)
        s.flush()
        full = _card_from_row(row)
    methods_mod.register(full)
    return full


# ------------------------------------------------------------------ demo reset (docx 21.2, zaxira reja)
def demo_reset():
    from . import seed as seed_mod
    from .models import Base
    with db.session() as s:
        for table in reversed(Base.metadata.sorted_tables):
            if table.name != "llm_calls":
                s.execute(table.delete())
    methods_mod.METHODS[:] = [m for m in methods_mod.METHODS if not m.get("custom")]
    for key in [k for k, v in methods_mod.METHOD_BY_ID.items() if v.get("custom")]:
        methods_mod.METHOD_BY_ID.pop(key)
    seed_mod.seed()
